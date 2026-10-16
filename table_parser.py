"""Buoc 1 - Chuan hoa bang tinh Excel thanh cac DataFrame "phang".

Nguyen ly:
  * Doc tung sheet bang openpyxl, dien gia tri o gop (merged cells) vao moi o ma no bao phu.
    O cong thuc khong co gia tri luu san -> tu tinh cac cong thuc don gian (=A1, SUM, AVERAGE, + - * /).
  * Tach sheet thanh cac khoi cot (2 bang dat canh nhau, ngan boi cot trong) va xu ly tung khoi.
  * Quet tung dong va phan loai: trong / tieu de (header) / du lieu / nhan nhom (section) / ghi chu.
    So luu dang chu ("1,234", "12.5%", "$30") cung tinh la du lieu.
  * Moi khoi header + cac dong du lieu phia duoi = 1 bang. Mot sheet co the co nhieu bang.
    Bang toan chu (khong co so): header = cac dong dau in dam / o gop, con lai la du lieu.
  * Ten cot = ghep cac tang tieu de bang " | " (vd "Some college or associate degree | Some college, no degree").
  * Dong chi co nhan (vd "TOTAL", "Women") la nhan nhom -> luu vao cot `section` cua cac dong ben duoi.
  * DataFrame: index = so dong Excel, cot = chu cai cot Excel (A, B, C...) + cot `section`.

Chay thu:  python table_parser.py duong_dan_file.xlsx   (in ra ban mo ta bang gui cho LLM)
"""
from __future__ import annotations

import datetime as dt
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

MAX_HEADER_ROWS = 6          # khoi tieu de dai hon -> phan du coi la du lieu
MAX_RENDER_ROWS = 400        # so dong toi da in ra cho LLM (df van giu du)
NOTE_PREFIX = re.compile(r"^(footnotes?|notes?|source|sources)\b|^\(\d+\)\s|^\d+/\s", re.I)
MISSING = {"", "-", "–", "—", "n/a", "na", "(na)", "(x)", "x", "..", "…"}

_FOOTNOTE = re.compile(r"(\s*\(\d+\))+\s*$")      # "Total(2)", "Hospital services(1)(3)"
_SLASH_NOTE = re.compile(r"\s+\d+/\s*$")           # "Trade balance 3/"
_NUM_STR = re.compile(r"^[-+]?[$€£]?\s*(\d{1,3}(,\d{3})+|\d+)(\.\d+)?\s*%?$|^[-+]?[$€£]?\.\d+%?$")
_DATE_STR = re.compile(r"^\d{4}-\d{1,2}-\d{1,2}([ T]\d{1,2}:\d{2}(:\d{2})?)?$")


# tien to giu nguyen dau gach noi khi bi ngat dong ("Self-\nemployed" -> "Self-employed")
_KEEP_HYPHEN = {"self", "non", "full", "part", "year", "high", "low", "well", "long", "short",
                "multi", "co", "ex", "semi", "anti", "pre", "post", "cross", "one", "two", "three",
                "four", "five", "half", "all", "first", "second", "third", "mid", "over", "under",
                "out", "in", "on", "off", "re", "e", "x", "sub", "inter", "intra", "single"}


def _join_hyphen(m) -> str:
    left, right = m.group(1), m.group(2)
    if left.lower() in _KEEP_HYPHEN or left[-1:].isdigit():
        return f"{left}-{right}"
    return f"{left}{right}"


def clean_text(s) -> str:
    """Lam sach chuoi: noi tu bi ngat gach noi, bo xuong dong, bo ky hieu chu thich."""
    if s is None:
        return ""
    s = str(s).replace("\r", "").replace("`", "'").replace("\xa0", " ")
    s = re.sub(r"(\w+)-\n\s*([a-z])", _join_hyphen, s)  # "Con-\nstruction" -> "Construction"
    s = re.sub(r"-\s*\n\s*", "-", s)                    # "2023-\nOct." -> "2023-Oct."
    s = re.sub(r"\s+", " ", s).strip()
    prev = None
    while prev != s:
        prev = s
        s = _SLASH_NOTE.sub("", _FOOTNOTE.sub("", s)).strip()
    return s


def _is_empty(v) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def _is_number(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _is_year(v) -> bool:
    return _is_number(v) and float(v).is_integer() and 1900 <= v <= 2100


def num_from_str(s):
    """'1,234' / '$5.2' / '12.5%' -> so (float); chuoi khac -> None. Giu nguyen ma co so 0 dau ('007')."""
    if not isinstance(s, str):
        return None
    t = s.strip()
    if not _NUM_STR.match(t) or re.match(r"^[-+]?0\d", t):
        return None
    try:
        return float(re.sub(r"[$€£%,\s+]", "", t))
    except ValueError:
        return None


def _is_data_value(v) -> bool:
    """Gia tri mang tinh 'du lieu' (khong phai nam trong tieu de)."""
    if isinstance(v, str):
        n = num_from_str(v)
        if n is not None:
            return not _is_year(n) or not float(n).is_integer() or "," in v or "%" in v
        return bool(_DATE_STR.match(v.strip()))
    return (_is_number(v) and not _is_year(v)) or isinstance(v, (dt.datetime, dt.date, dt.time))


@dataclass
class Table:
    name: str
    sheet: str
    title: str = ""
    captions: list = field(default_factory=list)
    columns: dict = field(default_factory=dict)      # chu cai cot -> ten cot day du
    df: pd.DataFrame | None = None
    notes: list = field(default_factory=list)
    lines: list = field(default_factory=list)        # (row, kind, payload) de in cho LLM


@dataclass
class _Grid:
    raw: list          # gia tri goc tung o
    filled: list       # gia tri sau khi dien o gop
    merges: list       # (r0, c0, r1, c1) chi so 0
    bold: list         # bold[r][c] = o co chu in dam
    ncols: int


# ----------------------------------------------------------------------------
def _load_grid(ws) -> _Grid:
    rows, bold = [], []
    for r in ws.iter_rows():
        rows.append([c.value for c in r])
        bold.append([bool(getattr(c, "font", None) is not None and c.font.b) for c in r])
    nrows = len(rows)
    ncols = max((len(r) for r in rows), default=0)
    for r, b in zip(rows, bold):
        r.extend([None] * (ncols - len(r)))
        b.extend([False] * (ncols - len(b)))
    for r in rows:
        for j, v in enumerate(r):
            if isinstance(v, str) and v.strip() == "":
                r[j] = None
    filled = [list(r) for r in rows]
    merges = []
    for mr in ws.merged_cells.ranges:
        r0, c0, r1, c1 = mr.min_row - 1, mr.min_col - 1, mr.max_row - 1, mr.max_col - 1
        if r0 >= nrows or c0 >= ncols:
            continue
        v = rows[r0][c0]
        merges.append((r0, c0, min(r1, nrows - 1), min(c1, ncols - 1)))
        for i in range(r0, min(r1, nrows - 1) + 1):
            for j in range(c0, min(c1, ncols - 1) + 1):
                filled[i][j] = v
    # bo cot trong o cuoi
    last = 0
    for r in rows:
        for j, v in enumerate(r):
            if not _is_empty(v):
                last = max(last, j + 1)
    return _Grid([r[:last] for r in rows], [r[:last] for r in filled], merges,
                 [b[:last] for b in bold], last)


def _row_kind(raw_row):
    vals = [(j, v) for j, v in enumerate(raw_row) if not _is_empty(v)]
    if not vals:
        return "blank"
    # "2 | Output per job: ..." = so thu tu bang + tieu de -> nhan, khong phai du lieu
    if len(vals) == 2 and vals[0][0] == 0 and _is_number(vals[0][1]) and float(vals[0][1]).is_integer() \
            and 0 < vals[0][1] < 1000 and isinstance(vals[1][1], str) and len(vals[1][1].split()) >= 3:
        return "label"
    data = [v for _, v in vals if _is_data_value(v)]
    # dong ma so trong tieu de: "Divisions | 10-12 | 13-15 | 28 | 29-30" (chi 1 so, con lai la ma dang chu)
    texts = [v for _, v in vals if isinstance(v, str)]
    if len(vals) >= 4 and len(data) == 1 and len(texts) == len(vals) - 1 \
            and sum(bool(re.fullmatch(r"\s*\d[\d,\-\s]*", v)) for v in texts) * 2 >= len(texts):
        return "header"
    if data:
        return "data"
    if len(vals) == 1 and vals[0][0] == 0:
        return "label"           # chi co text o cot dau: section / tieu de / ghi chu
    return "header"


def _column_blocks(g: _Grid) -> list[tuple[int, int]]:
    """Tach cac bang dat canh nhau (ngan boi cot hoan toan trong) thanh cac khoi cot [c0, c1)."""
    empty = [all(_is_empty(r[j]) for r in g.raw) for j in range(g.ncols)]
    blocks, start = [], None
    for j in range(g.ncols + 1):
        if j < g.ncols and not empty[j]:
            start = j if start is None else start
        elif start is not None:
            blocks.append((start, j))
            start = None
    if len(blocks) <= 1:
        return [(0, g.ncols)]

    def nonempty(r):
        return sum(not _is_empty(v) for v in g.raw[r])

    def first_row(b):
        return next((i for i, r in enumerate(g.raw) if any(not _is_empty(v) for v in r[b[0]:b[1]])), -1)

    def own_labels(b):
        # cot dau cua khoi co nhan chu tren cac dong co du lieu so -> bang rieng
        for r in g.raw:
            part = r[b[0]:b[1]]
            if isinstance(part[0], str) and num_from_str(part[0]) is None \
                    and any(_is_data_value(v) for v in part[1:]):
                return True
        return False

    out = [blocks[0]]
    for b in blocks[1:]:
        prev = out[-1]
        crossing = any(c0 < prev[1] and c1 >= b[0] and nonempty(r0) > 1 for r0, c0, r1, c1 in g.merges)
        separate = own_labels(b) or first_row(b) != first_row(prev)
        if crossing or not separate or b[1] - b[0] < 2 or prev[1] - prev[0] < 2:
            out[-1] = (prev[0], b[1])
        else:
            out.append(b)
    if out[0][0] > 0:
        out[0] = (0, out[0][1])
    return out


def parse_sheet(ws, sheet_index: int, table_offset: int) -> list[Table]:
    g = _load_grid(ws)
    if g.ncols == 0:
        return []
    tables: list[Table] = []
    for c0, c1 in _column_blocks(g):
        sub = _Grid([r[c0:c1] for r in g.raw], [r[c0:c1] for r in g.filled],
                    [(r0, max(a, c0) - c0, r1, min(b, c1 - 1) - c0)
                     for r0, a, r1, b in g.merges if a < c1 and b >= c0],
                    [r[c0:c1] for r in g.bold], c1 - c0)
        # khoi bat dau sau cot A: bo cac cot trong phia truoc de nhan nhom nam o cot 0
        lead = 0
        while lead < sub.ncols - 1 and all(_is_empty(r[lead]) for r in sub.raw):
            lead += 1
        if lead:
            sub = _Grid([r[lead:] for r in sub.raw], [r[lead:] for r in sub.filled],
                        [(r0, max(a - lead, 0), r1, b - lead) for r0, a, r1, b in sub.merges if b >= lead],
                        [r[lead:] for r in sub.bold], sub.ncols - lead)
        tables += _parse_block(sub, ws.title, sheet_index, table_offset + len(tables), c0 + lead)
    return tables


def _header_depth(hdr, g: _Grid) -> int:
    """So dong header thuc su cua 1 khoi dong toan chu: cac dong o gop ngang / in dam o dau khoi."""
    h0, last = hdr[0], hdr[-1]
    d = 1
    for r0, c0, r1, c1 in g.merges:
        if h0 <= r0 <= min(last, h0 + 2) and (r0 == h0 or c1 > c0):
            d = max(d, min(r1, last) - h0 + 1)
    nb = 0
    for r in hdr:
        cells = [b for v, b in zip(g.raw[r], g.bold[r]) if not _is_empty(v)]
        if cells and all(cells):
            nb += 1
        else:
            break
    if 0 < nb < len(hdr):
        d = max(d, nb)
    return min(d, len(hdr))


def _parse_block(g: _Grid, sheet, sheet_index, table_offset, off) -> list[Table]:
    raw, filled, merges, ncols = g.raw, g.filled, g.merges, g.ncols
    if ncols == 0:
        return []
    kinds = [_row_kind(r) for r in raw]
    # o gop o cot dau trai dai nhieu dong (khong full chieu ngang) -> chac chan la khoi header
    a_merge_header = set()
    for r0, c0, r1, c1 in merges:
        if c0 == 0 and r1 > r0 and c1 < ncols - 1:
            for i in range(r0, r1 + 1):
                if i < len(kinds) and kinds[i] != "data":
                    a_merge_header.add(i)
                    if kinds[i] in ("blank", "label"):
                        kinds[i] = "header"

    tables: list[Table] = []
    pre_text: list[str] = []
    i, n = 0, len(raw)
    while i < n:
        if kinds[i] != "header":
            if kinds[i] == "label":
                pre_text.append(" ".join(clean_text(v) for v in raw[i] if not _is_empty(v)))
            if kinds[i] == "data" and not tables:
                # bang khong co header -> dung ten cot mac dinh
                i = _consume_table(g, kinds, sheet, [], i, tables, pre_text, table_offset, off)
                pre_text = []
                continue
            i += 1
            continue
        # khoi header
        h0 = i
        while i < n and (kinds[i] == "header" or (kinds[i] == "blank" and i in a_merge_header)):
            i += 1
        hdr = list(range(h0, i))
        # dong tieu de / chu thich hep (1-2 o, khong gop ngang) o dau khoi header rong -> la text truoc bang
        width = max(sum(not _is_empty(v) for v in raw[k]) for k in hdr)
        hmerge = {r0 for r0, c0, r1, c1 in merges if c1 > c0}
        while len(hdr) > 1 and width >= 5 and hdr[0] not in hmerge \
                and sum(not _is_empty(v) for v in raw[hdr[0]]) <= 2:
            pre_text.append(" ".join(clean_text(v) for v in raw[hdr[0]] if not _is_empty(v)))
            hdr.pop(0)
        j = i
        while j < n and kinds[j] == "blank":
            j += 1
        followed = j < n and kinds[j] in ("data", "label")
        text_rows = [k for k in hdr if kinds[k] == "header"]
        if (not followed and len(text_rows) >= 2) or len(hdr) > MAX_HEADER_ROWS:
            # bang toan chu (hoac khoi header qua dai): chi giu cac dong header that, con lai la du lieu
            d = _header_depth(hdr, g)
            if d < len(hdr):
                for k in hdr[d:]:
                    if kinds[k] == "header":
                        kinds[k] = "data"
                hdr = hdr[:d]
                i = j = hdr[-1] + 1
                while j < n and kinds[j] == "blank":
                    j += 1
                followed = j < n and kinds[j] in ("data", "label")
        if not followed:
            # header khong co du lieu ben duoi -> coi la text
            for k in hdr:
                pre_text.append(" ".join(clean_text(v) for v in raw[k] if not _is_empty(v)))
            continue
        i = _consume_table(g, kinds, sheet, hdr, j, tables, pre_text, table_offset, off)
        pre_text = []
    return tables


def _wrapped(prev: str, cur: str) -> bool:
    """`cur` la phan tiep theo cua `prev` bi xuong dong (khong phai cap tieu de moi)."""
    return bool(re.search(r"[,&]$|[a-z]-$", prev)) and not re.match(r"\d", cur) \
        or (cur.startswith("&") and not prev.endswith("&"))


def _consume_table(g: _Grid, kinds, sheet, hdr, start, tables, pre_text, table_offset, off):
    raw, filled, ncols = g.raw, g.filled, g.ncols
    n = len(raw)
    letter = [get_column_letter(c + 1 + off) for c in range(ncols)]
    t = Table(name=f"tables[{table_offset + len(tables)}]", sheet=sheet)
    title_parts = [p for p in pre_text if p]
    # --- ten cot -------------------------------------------------------------
    caption_rows = set()
    for r in hdr:
        tail = filled[r][1:]
        if tail and all(not _is_empty(v) for v in tail) and len({str(v) for v in tail}) == 1 and ncols > 2:
            caption_rows.add(r)
            t.captions.append(clean_text(tail[0]))
    for c in range(ncols):
        parts = []
        for r in hdr:
            if r in caption_rows and c > 0:
                continue
            v = filled[r][c]
            if _is_empty(v):
                continue
            if isinstance(v, (dt.datetime, dt.date)):
                s = v.strftime("%Y-%m-%d")
            elif _is_number(v):
                s = str(int(v)) if float(v).is_integer() else str(v)
            else:
                s = clean_text(v)
            if s and parts and r - 1 in hdr and filled[r - 1][c] is raw[r - 1][c] and _wrapped(parts[-1], s):
                # chu bi xuong dong trong o tieu de: "Rubber," / "plastics &" / "Pharmaceutic-" "als"
                parts[-1] = parts[-1][:-1] + s if parts[-1].endswith("-") else parts[-1] + " " + s
            elif s and (not parts or parts[-1] != s):
                parts.append(s)
        t.columns[letter[c]] = " | ".join(parts) if parts else f"(column {letter[c]})"
    # tieu de: dong text truoc bang + caption cot A neu co
    t.title = " ".join(title_parts)
    # --- du lieu -------------------------------------------------------------
    # o gop doc trong vung du lieu (vd "Region" gop nhieu dong) -> moi dong deu mang gia tri
    vfill = {(i, c) for r0, c0, r1, c1 in g.merges if r1 > r0 and r0 >= start
             for i in range(r0, r1 + 1) for c in range(c0, c1 + 1)}
    records, index = [], []
    root, sub = None, None
    root_has_data = False
    i = start
    in_notes = False
    while i < n:
        k = kinds[i]
        if k == "header":
            # header moi sau 1 dong trong -> bang moi
            if i > 0 and kinds[i - 1] == "blank":
                break
            k = "data"   # dong text nhieu cot trong vung du lieu -> du lieu dang chu
        if k == "blank":
            i += 1
            continue
        if k == "label":
            txt = " ".join(str(v).strip() for v in raw[i] if not _is_empty(v))
            if in_notes or NOTE_PREFIX.search(txt) or len(txt) > 90:
                in_notes = True
                t.notes.append(clean_text(txt))
                t.lines.append((i + 1, "note", clean_text(txt)))
            else:
                lab = clean_text(txt)
                if lab.isupper():
                    root, sub = lab, None
                    root_has_data = False
                else:
                    if root and root_has_data and sub is None:
                        root = None      # vd "TOTAL" co du lieu rieng -> "Men" la nhom ngang hang
                    sub = lab
                t.lines.append((i + 1, "section", lab))
            i += 1
            continue
        if in_notes:
            break
        row = {}
        for c in range(ncols):
            v = filled[i][c] if (i, c) in vfill else raw[i][c]
            if isinstance(v, str):
                v = None if v.strip().lower() in MISSING else clean_text(v)
            elif isinstance(v, (dt.datetime, dt.date)):
                v = pd.Timestamp(v)
            row[letter[c]] = v
        row["section"] = " > ".join(p for p in (root, sub) if p)
        root_has_data = True
        records.append(row)
        index.append(i + 1)
        t.lines.append((i + 1, "data", row))
        i += 1
    cols = list(t.columns) + ["section"]
    df = pd.DataFrame(records, index=pd.Index(index, name="excel_row"), columns=cols)
    for c in list(t.columns):
        if t.columns[c].startswith("(column ") and df[c].isna().all():
            del t.columns[c]             # cot hoan toan trong (thuoc bang ben canh)
            df = df.drop(columns=c)
            continue
        df[c] = _convert_column(df[c])
    t.df = df
    if len(df):
        tables.append(t)
    return i


def _convert_column(s: pd.Series) -> pd.Series:
    """Ep kieu cot: ngay thang -> datetime, so (ke ca so dang chu '1,234', '12%') -> so."""
    nonnull = s.dropna()
    if nonnull.empty:
        return s
    if nonnull.map(lambda v: isinstance(v, (pd.Timestamp, dt.date))).any():
        return pd.to_datetime(s, errors="coerce")
    if nonnull.map(lambda v: isinstance(v, str) and bool(_DATE_STR.match(v))).all():
        return pd.to_datetime(s, errors="coerce")
    conv = s.map(lambda v: num_from_str(v) if isinstance(v, str) else v)
    num = pd.to_numeric(conv, errors="coerce")
    if num.notna().sum() == len(nonnull):
        return num
    return s


# ----------------------------------------------------------------------------
# Tinh cong thuc don gian khi file khong luu san gia tri (vd file tao bang code)
_RANGE_RE = re.compile(r"(?:(?:'([^']+)'|([A-Za-z_][\w.]*))!)?\$?([A-Z]{1,3})\$?(\d+):\$?([A-Z]{1,3})\$?(\d+)")
_CELL_RE = re.compile(r"(?:(?:'([^']+)'|([A-Za-z_][\w.]*))!)?\$?([A-Z]{1,3})\$?(\d+)(?![\w(])")
_FUNC_RE = re.compile(r"\b(SUM|AVERAGE|MIN|MAX|COUNT|ABS|ROUND)\(")


def _flat(args):
    out = []
    for a in args:
        if isinstance(a, list):
            out += _flat(a)
        elif _is_number(a):
            out.append(a)
    return out


_XL_FUNCS = {
    "_SUM": lambda *a: sum(_flat(a)),
    "_AVERAGE": lambda *a: sum(_flat(a)) / len(_flat(a)),
    "_MIN": lambda *a: min(_flat(a)),
    "_MAX": lambda *a: max(_flat(a)),
    "_COUNT": lambda *a: len(_flat(a)),
    "_ABS": abs,
    "_ROUND": lambda x, d=0: round(x, int(d)),
}


def _eval_formula(formula: str, sheet: str, get):
    expr = formula[1:].strip()

    def rng(m):
        sh = m.group(1) or m.group(2) or sheet
        c0, c1 = column_index_from_string(m.group(3)), column_index_from_string(m.group(5))
        r0, r1 = int(m.group(4)), int(m.group(6))
        return repr([get(sh, r, c) for r in range(r0, r1 + 1) for c in range(c0, c1 + 1)])

    def cell(m):
        v = get(m.group(1) or m.group(2) or sheet, int(m.group(4)), column_index_from_string(m.group(3)))
        if isinstance(v, (dt.datetime, dt.date)):
            raise ValueError("date in formula")
        return repr(0 if v is None else v)

    expr = _FUNC_RE.sub(lambda m: f"_{m.group(1)}(", expr)
    expr = _RANGE_RE.sub(rng, expr)
    expr = _CELL_RE.sub(cell, expr).replace("^", "**")
    return eval(expr, {"__builtins__": {}}, dict(_XL_FUNCS))  # noqa: S307 - chi so va ham tren


def _fill_formulas(path, wb) -> None:
    """Dien gia tri cho o cong thuc ma file khong luu ket qua (data_only tra ve None)."""
    try:
        fwb = load_workbook(path, data_only=False)
    except Exception:  # noqa: BLE001
        return
    todo = []
    for fws in fwb.worksheets:
        if fws.title not in wb.sheetnames:
            continue
        ws = wb[fws.title]
        for row in fws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and ws[c.coordinate].value is None:
                    todo.append((ws, c.coordinate, c.value))
    if not todo:
        return

    def get(sh, r, c):
        return wb[sh].cell(row=r, column=c).value if sh in wb.sheetnames else None

    for _ in range(5):                      # vai vong cho cong thuc phu thuoc nhau
        left = []
        for ws, coord, f in todo:
            try:
                v = _eval_formula(f, ws.title, get)
                ws[coord].value = v
            except Exception:  # noqa: BLE001
                left.append((ws, coord, f))
        if len(left) == len(todo):
            break
        todo = left


# ----------------------------------------------------------------------------
def parse_workbook(path) -> list[Table]:
    path = Path(path)
    if path.suffix.lower() in (".csv", ".tsv"):
        sep = "\t" if path.suffix.lower() == ".tsv" else ","
        frame = pd.read_csv(path, sep=sep, header=None, dtype=object)
        return _parse_frame(frame, path.stem)
    if path.suffix.lower() == ".xls":
        tables = []
        for sh, frame in pd.read_excel(path, sheet_name=None, header=None).items():
            tables += _parse_frame(frame, sh, len(tables))
        return tables
    wb = load_workbook(path, data_only=True)
    _fill_formulas(path, wb)
    tables: list[Table] = []
    for si, ws in enumerate(wb.worksheets):
        if ws.sheet_state != "visible" and len(wb.worksheets) > 1:
            continue
        try:
            tables += parse_sheet(ws, si, len(tables))
        except Exception as e:  # noqa: BLE001 - van tra ve bang tho de LLM tu doc
            print(f"[warn] parse_sheet '{ws.title}' failed ({e}); using raw grid")
            tables += _raw_table(ws, len(tables))
    return tables


def _raw_table(ws, offset) -> list[Table]:
    """Du phong: ca sheet thanh 1 bang, khong header."""
    g = _load_grid(ws)
    kinds = ["blank" if all(_is_empty(v) for v in r) else "data" for r in g.raw]
    out: list[Table] = []
    start = next((i for i, k in enumerate(kinds) if k == "data"), None)
    if start is not None:
        _consume_table(g, kinds, ws.title, [], start, out, [], offset, 0)
    return out


def _parse_frame(frame, sheet, offset=0):
    """Ho tro csv/xls: dung lai logic bang cach tao workbook tam trong bo nho."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = str(sheet)[:31]
    for row in frame.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else _num(v) for v in row])
    return parse_sheet(ws, 0, offset)


def _num(v):
    if isinstance(v, str):
        f = num_from_str(v)
        if f is not None and "%" not in v:
            return int(f) if f.is_integer() else f
    return v


# ----------------------------------------------------------------------------
def _fmt(v) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)) or v is pd.NaT:
        return ""
    if isinstance(v, pd.Timestamp):
        return v.strftime("%Y-%m-%d") if v.time() == dt.time(0) else str(v)
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(round(v, 10))
    return str(v)


def _profile(t: Table, max_values: int = 30) -> list[str]:
    """Tom tat tung cot (dung khi bang qua dai, khong in het duoc): kieu, khoang gia tri, cac gia tri chu."""
    out = ["Column summary (computed over ALL rows):"]
    for c in t.columns:
        s = t.df[c].dropna()
        if s.empty:
            out.append(f"  {c}: empty")
        elif pd.api.types.is_datetime64_any_dtype(s):
            out.append(f"  {c}: date {_fmt(s.min())} .. {_fmt(s.max())}")
        elif pd.api.types.is_numeric_dtype(s):
            out.append(f"  {c}: number {_fmt(s.min())} .. {_fmt(s.max())}, {len(s)} values")
        else:
            vc = s.astype(str).value_counts()
            if len(vc) > max_values and len(vc) > len(s) / 2:     # gan nhu duy nhat (ten, ma...)
                ex = ", ".join(f"{v!r}" for v in s.astype(str).head(3))
                out.append(f"  {c}: text, {len(vc)} distinct (e.g. {ex})")
                continue
            vals = ", ".join(f"{v!r}" for v in vc.index[:max_values])
            more = f", ... (+{len(vc) - max_values} more)" if len(vc) > max_values else ""
            out.append(f"  {c}: text, {len(vc)} distinct: {vals}{more}")
    return out


def render_tables(tables: list[Table], max_rows: int = MAX_RENDER_ROWS, head_rows: int = 60) -> str:
    """Ban mo ta bang dang text gon de gui cho LLM.
    Bang dai hon max_rows: chi in head_rows dong dau + tom tat tung cot (tiet kiem token)."""
    out = []
    for t in tables:
        n_data = sum(1 for _, kind, _ in t.lines if kind == "data")
        limit = max_rows if n_data <= max_rows else head_rows
        out.append(f"### {t.name}  (sheet '{t.sheet}', {len(t.df)} data rows)")
        if t.title:
            out.append(f"Title: {t.title[:400]}")
        if t.captions:
            out.append("Info: " + "; ".join(t.captions))
        out.append("Columns:")
        for letter, name in t.columns.items():
            out.append(f"  {letter} = {name}")
        out.append("Rows  (Excel row: " + " | ".join(t.columns) + "),  [..] = section heading, '' = empty:")
        shown = 0
        for row, kind, payload in t.lines:
            if kind == "section":
                if shown < limit:
                    out.append(f"[{payload}]")
            elif kind == "note":
                continue
            else:
                shown += 1
                if shown > limit:
                    continue
                out.append(f"{row}: " + " | ".join(_fmt(payload[c]) for c in t.columns))
        if shown > limit:
            out.append(f"... ({shown - limit} more rows not shown; they ARE in the DataFrame)")
            out += _profile(t)
        if t.notes:
            out.append("Notes: " + " ".join(t.notes)[:600])
        out.append("")
    return "\n".join(out)


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    for p in sys.argv[1:]:
        tabs = parse_workbook(p)
        print(f"===== {p}: {len(tabs)} table(s)")
        print(render_tables(tabs))
        for t in tabs:
            print(t.df.head(3).to_string())
